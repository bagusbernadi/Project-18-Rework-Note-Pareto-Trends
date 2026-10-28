"""10: Cek silang Excel (Q14). Pareto label klasifikasi vs label manual pada golden set dihitung dengan formula Excel (COUNTIF, RANK),
dihitung ulang oleh LibreOffice (headless), lalu dibandingkan dengan view SQL. Workbook dengan nilai terhitung: outputs/crosscheck.xlsx."""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import openpyxl
from openpyxl.styles import Font
import pandas as pd
import common as CM

g = pd.read_csv(CM.RES2 / "compare_golden.csv").sort_values("note_id")
sql = pd.read_csv(CM.RES2 / "pareto_golden.csv").set_index("category_code")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "golden"
ws.append(["note_id", "label_manual", "label_klasifikasi"])
for r in g.itertuples():
    ws.append([r.note_id, r.manual_cat, r.llm_cat])
n = len(g) + 1
pr = wb.create_sheet("pareto")
pr.append(["kategori", "n_manual", "n_klasifikasi", "rank_manual", "rank_klasifikasi"])
for c in pr[1]:
    c.font = Font(bold=True)
cats = list(CM.CATS)
last = len(cats) + 1
for i, c in enumerate(cats, start=2):
    # peringkat dengan pemecah seri menurut urutan kode: RANK + jumlah kategori sebelumnya yang nilainya sama
    pr.append([c, f"=COUNTIF(golden!B2:B{n},A{i})", f"=COUNTIF(golden!C2:C{n},A{i})",
               f"=RANK(B{i},B$2:B${last},0)+COUNTIF(B$2:B{i - 1},B{i})" if i > 2 else f"=RANK(B{i},B$2:B${last},0)",
               f"=RANK(C{i},C$2:C${last},0)+COUNTIF(C$2:C{i - 1},C{i})" if i > 2 else f"=RANK(C{i},C$2:C${last},0)"])
tmp = Path(tempfile.mkdtemp())
src = tmp / "crosscheck_raw.xlsx"
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)], check=True, capture_output=True, timeout=240)
calc = tmp / "out" / src.name
shutil.copy(calc, CM.OUTPUTS / "crosscheck.xlsx")
vals = openpyxl.load_workbook(calc, data_only=True)["pareto"]
rows = []
for i, c in enumerate(cats, start=2):
    rows.append(dict(category_code=c, n_manual_excel=vals.cell(i, 2).value, n_manual_sql=int(sql.loc[c, "n_manual"]),
                     n_klasifikasi_excel=vals.cell(i, 3).value, n_klasifikasi_sql=int(sql.loc[c, "n_llm"]),
                     rank_manual_excel=vals.cell(i, 4).value, rank_manual_sql=int(sql.loc[c, "rank_manual"]),
                     rank_klasifikasi_excel=vals.cell(i, 5).value, rank_klasifikasi_sql=int(sql.loc[c, "rank_llm"])))
res = pd.DataFrame(rows)
res["cocok"] = ((res.n_manual_excel == res.n_manual_sql) & (res.n_klasifikasi_excel == res.n_klasifikasi_sql)
                & (res.rank_manual_excel == res.rank_manual_sql) & (res.rank_klasifikasi_excel == res.rank_klasifikasi_sql)).astype(int)
res.to_csv(CM.RES2 / "crosscheck_excel.csv", index=False)
print(res.to_string(index=False))
assert res.cocok.sum() == len(res), "Excel dan SQL tidak cocok"
